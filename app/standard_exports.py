"""Fixed-column lead and sales-deal exports required by the motivation brief."""
from __future__ import annotations

import asyncio
from datetime import date, datetime, time, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import store
from app.bitrix_client import BitrixClient
from app.config import get_settings


def _name_map(items: list[dict]) -> dict[str, str]:
    return {
        str(item.get("STATUS_ID") or ""): str(item.get("NAME") or item.get("STATUS_ID") or "")
        for item in items
    }


def _first_phone(item: dict) -> str:
    for phone in item.get("PHONE") or []:
        value = str((phone or {}).get("VALUE") or "").strip()
        if value:
            return value
    return ""


def _handoff_at(entity_type: str, entity_id: str) -> str:
    rows = store.rows(
        """SELECT occurred_at FROM crm_item_events
           WHERE entity_type=? AND entity_id=? AND event_type IN ('assignee_changed','transferred')
           ORDER BY occurred_at ASC,id ASC LIMIT 1""",
        (entity_type, entity_id),
    )
    return str(rows[0]["occurred_at"]) if rows else ""


async def collect_standard_sales_export(
    start: date, end: date, *, client: BitrixClient | None = None,
) -> dict:
    settings = get_settings()
    client = client or BitrixClient()
    timezone = ZoneInfo(settings.daily_stats_timezone)
    start_iso = datetime.combine(start, time.min, timezone).isoformat()
    end_iso = datetime.combine(end + timedelta(days=1), time.min, timezone).isoformat()
    reason_field = settings.deal_refusal_reason_field
    leads, created_deals, modified_deals, sources, lead_stages, deal_stages, fields = await asyncio.gather(
        client.get_leads_created_between_full(start_iso, end_iso),
        client.get_deals_created_between_full(
            settings.track_deal_category_id, start_iso, end_iso, extra_fields=[reason_field],
        ),
        client.get_deals_modified_between(settings.track_deal_category_id, start_iso, end_iso),
        client.get_sources(), client.get_lead_statuses(),
        client.get_deal_stages(settings.track_deal_category_id), client.get_deal_userfields(),
    )
    # Modified-deal query uses the common select, so load full rows for any IDs not in creation cohort.
    deals_by_id = {str(item["ID"]): item for item in created_deals}
    missing_ids = [str(item["ID"]) for item in modified_deals if str(item["ID"]) not in deals_by_id]
    for item in await client.get_deals_by_ids(
        missing_ids, select=list(dict.fromkeys([
            "ID", "DATE_CREATE", "DATE_MODIFY", "ASSIGNED_BY_ID", "STAGE_ID", "SOURCE_ID",
            "CONTACT_ID", "OPPORTUNITY", "CURRENCY_ID", reason_field,
        ])),
    ):
        deals_by_id[str(item["ID"])] = item
    deals = list(deals_by_id.values())
    contact_ids = [str(item.get("CONTACT_ID") or "") for item in deals]
    contacts, history = await asyncio.gather(
        client.get_contacts_full_by_ids(contact_ids),
        client.get_deal_stage_history([str(item["ID"]) for item in deals]),
    )
    contacts_by_id = {str(item["ID"]): item for item in contacts}
    source_names = _name_map(sources)
    lead_stage_names = _name_map(lead_stages)
    deal_stage_names = _name_map(deal_stages)
    manager_ids = sorted({
        str(item.get("ASSIGNED_BY_ID") or "") for item in [*leads, *deals]
        if item.get("ASSIGNED_BY_ID")
    })
    manager_values = await asyncio.gather(*(client.get_user_name(user_id) for user_id in manager_ids))
    manager_names = {
        user_id: name or f"Сотрудник #{user_id}"
        for user_id, name in zip(manager_ids, manager_values)
    }
    reason_values: dict[str, str] = {}
    reason_definition = next(
        (item for item in fields if str(item.get("FIELD_NAME") or "") == reason_field), None,
    )
    if reason_definition:
        reason_values = {
            str(item.get("ID")): str(item.get("VALUE") or "")
            for item in reason_definition.get("LIST") or []
        }
    last_stage_at: dict[str, str] = {}
    for event in history:
        deal_id = str(event.get("OWNER_ID") or "")
        if str(event.get("CATEGORY_ID") or "0") == str(settings.track_deal_category_id):
            last_stage_at[deal_id] = str(event.get("CREATED_TIME") or "")
    lead_rows = []
    for lead in leads:
        lead_id = str(lead["ID"])
        manager_id = str(lead.get("ASSIGNED_BY_ID") or "")
        source_id = str(lead.get("SOURCE_ID") or "")
        stage_id = str(lead.get("STATUS_ID") or "")
        lead_rows.append({
            "id": lead_id, "created_at": str(lead.get("DATE_CREATE") or ""),
            "handoff_at": _handoff_at("lead", lead_id),
            "manager": manager_names.get(manager_id, "Без ответственного"),
            "stage": lead_stage_names.get(stage_id, stage_id),
            "source": source_names.get(source_id, source_id or "Без источника"),
            "phone": _first_phone(lead),
        })
    deal_rows = []
    for deal in deals:
        deal_id = str(deal["ID"])
        manager_id = str(deal.get("ASSIGNED_BY_ID") or "")
        source_id = str(deal.get("SOURCE_ID") or "")
        stage_id = str(deal.get("STAGE_ID") or "")
        reason_id = str(deal.get(reason_field) or "")
        contact = contacts_by_id.get(str(deal.get("CONTACT_ID") or ""), {})
        deal_rows.append({
            "id": deal_id, "created_at": str(deal.get("DATE_CREATE") or ""),
            "manager": manager_names.get(manager_id, "Без ответственного"),
            "stage": deal_stage_names.get(stage_id, stage_id),
            "stage_changed_at": last_stage_at.get(deal_id, str(deal.get("DATE_MODIFY") or "")),
            "source": source_names.get(source_id, source_id or "Без источника"),
            "phone": _first_phone(contact), "amount": deal.get("OPPORTUNITY") or 0,
            "currency": str(deal.get("CURRENCY_ID") or ""),
            "refusal_reason": reason_values.get(reason_id, reason_id),
        })
    return {"start": start.isoformat(), "end": end.isoformat(), "leads": lead_rows, "deals": deal_rows}


def _excel_text(value: object) -> str:
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _header(sheet, headers: list[str]) -> None:
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill; cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_standard_sales_workbook(data: dict) -> bytes:
    workbook = Workbook()
    leads = workbook.active; leads.title = "Лиды"
    _header(leads, [
        "ID", "Дата создания", "Передан менеджеру", "Ответственный",
        "Стадия", "Источник", "Телефон",
    ])
    for item in data["leads"]:
        leads.append([
            item["id"], item["created_at"], item["handoff_at"], item["manager"],
            _excel_text(item["stage"]), _excel_text(item["source"]), _excel_text(item["phone"]),
        ])
        leads.cell(leads.max_row, 7).number_format = "@"
    deals = workbook.create_sheet("Сделки")
    _header(deals, [
        "ID", "Дата создания", "Ответственный", "Стадия", "Дата изменения стадии",
        "Источник", "Телефон контакта", "Сумма", "Валюта", "Причина отказа",
    ])
    for item in data["deals"]:
        deals.append([
            item["id"], item["created_at"], item["manager"], _excel_text(item["stage"]),
            item["stage_changed_at"], _excel_text(item["source"]), _excel_text(item["phone"]),
            item["amount"], item["currency"], _excel_text(item["refusal_reason"]),
        ])
        deals.cell(deals.max_row, 7).number_format = "@"
    for sheet, widths in ((leads, (12, 22, 22, 25, 28, 26, 20)), (deals, (12, 22, 25, 28, 22, 26, 20, 16, 12, 28))):
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO(); workbook.save(buffer); return buffer.getvalue()
