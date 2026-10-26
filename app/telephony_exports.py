"""Fixed-column MegaFon exports used by sales motivation reports."""
from __future__ import annotations

import asyncio
from datetime import date, datetime, time, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import store
from app.bitrix_client import BitrixClient
from app.config import get_settings


def _direction(value: object) -> str:
    raw = str(value or "").casefold()
    if raw in {"out", "outbound", "outgoing"}:
        return "Исходящий"
    if raw in {"in", "inbound", "incoming"}:
        return "Входящий"
    return str(value or "Не определено")


async def collect_telephony_export(
    start: date, end: date, *, client: BitrixClient | None = None,
) -> dict:
    settings = get_settings()
    client = client or BitrixClient()
    timezone = ZoneInfo(settings.daily_stats_timezone)
    start_at = datetime.combine(start, time.min, timezone)
    end_at = datetime.combine(end + timedelta(days=1), time.min, timezone)
    rows = store.rows(
        """SELECT * FROM megafon_calls
           WHERE started_at>=? AND started_at<? ORDER BY started_at,call_id""",
        (start_at.astimezone(ZoneInfo("UTC")).isoformat(), end_at.astimezone(ZoneInfo("UTC")).isoformat()),
    )
    user_map = getattr(settings, "megafon_bitrix_users", {})
    user_ids = sorted({str(user_map.get(str(row.get("manager_extension") or "")) or "") for row in rows} - {""})
    names = await asyncio.gather(*(client.get_user_name(user_id) for user_id in user_ids))
    names_by_id = {user_id: name or f"Сотрудник #{user_id}" for user_id, name in zip(user_ids, names)}
    items = []
    for row in rows:
        started = datetime.fromisoformat(str(row["started_at"]).replace("Z", "+00:00")).astimezone(timezone)
        extension = str(row.get("manager_extension") or "")
        user_id = str(user_map.get(extension) or "")
        status = str(row.get("call_status") or "")
        answered = status.casefold() in {"success", "answered", "connected"}
        items.append({
            "call_id": str(row.get("call_id") or ""),
            "type": _direction(row.get("direction")),
            "client": str(row.get("client_phone") or ""),
            "phone": str(row.get("client_phone") or ""),
            "employee": names_by_id.get(user_id, f"Внутренний номер {extension}" if extension else "Не определён"),
            "date": started.date().isoformat(), "time": started.strftime("%H:%M:%S"),
            "started_at": started.isoformat(), "wait_seconds": int(row.get("wait_seconds") or 0),
            "duration_seconds": int(row.get("duration_seconds") or 0),
            "direction": _direction(row.get("direction")), "answered": answered,
            "dial_result": "Дозвон" if answered else "Не дозвонился",
        })
    return {"start": start.isoformat(), "end": end.isoformat(), "items": items}


def _duration(seconds: int) -> str:
    hours, remainder = divmod(max(0, seconds), 3600)
    minutes, rest = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{rest:02d}"


def _header(sheet, headers: list[str]) -> None:
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill; cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_telephony_workbook(data: dict) -> bytes:
    workbook = Workbook()
    history = workbook.active; history.title = "Телефония АТС"
    _header(history, ["Тип звонка", "Клиент", "Сотрудник", "Дата", "Время", "Ожидание", "Длительность"])
    mine = workbook.create_sheet("Мои звонки")
    _header(mine, ["Дата", "Сотрудник", "Номер", "Направление", "Отвечен", "Длительность", "Дозвон"])
    for item in data["items"]:
        history.append([
            item["type"], "'" + item["client"], item["employee"],
            date.fromisoformat(item["date"]), item["time"], _duration(item["wait_seconds"]),
            _duration(item["duration_seconds"]),
        ])
        mine.append([
            date.fromisoformat(item["date"]), item["employee"], "'" + item["phone"],
            item["direction"], "Да" if item["answered"] else "Нет",
            _duration(item["duration_seconds"]), item["dial_result"],
        ])
        history.cell(history.max_row, 2).number_format = "@"
        mine.cell(mine.max_row, 3).number_format = "@"
        history.cell(history.max_row, 4).number_format = "DD.MM.YYYY"
        mine.cell(mine.max_row, 1).number_format = "DD.MM.YYYY"
    for sheet, widths in ((history, (16, 20, 25, 14, 12, 14, 14)), (mine, (14, 25, 20, 16, 12, 14, 16))):
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO(); workbook.save(buffer); return buffer.getvalue()
