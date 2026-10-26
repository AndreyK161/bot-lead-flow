"""Manager workday/break history from the LeadFlow presence collector."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import get_settings


def _parse_datetime(value: object) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None


def _seconds(start: datetime, end: datetime) -> int:
    return max(0, int((end - start).total_seconds()))


def build_worktime_payload(
    raw: dict,
    start: date,
    end: date,
    *,
    now: datetime | None = None,
    timezone_name: str = "Europe/Moscow",
) -> dict:
    timezone = ZoneInfo(timezone_name)
    current = (now or datetime.now(timezone)).astimezone(timezone)
    items: list[dict] = []
    for manager in raw.get("managers") or []:
        events_by_day: dict[date, list[tuple[str, datetime]]] = {}
        for event in manager.get("events") or []:
            occurred_at = _parse_datetime(event.get("occurred_at"))
            if not occurred_at:
                continue
            local = occurred_at.astimezone(timezone)
            events_by_day.setdefault(local.date(), []).append((str(event.get("event_type")), local))
        days = set(events_by_day)
        if start <= current.date() <= end:
            days.add(current.date())
        for day in sorted(days):
            if not start <= day <= end:
                continue
            day_events = sorted(events_by_day.get(day, []), key=lambda item: item[1])
            shift_start = next((moment for kind, moment in day_events if kind == "workday_started"), None)
            shift_end = next((moment for kind, moment in reversed(day_events) if kind == "workday_ended"), None)
            breaks: list[dict] = []
            open_break: datetime | None = None
            for kind, moment in day_events:
                if kind == "lunch_started":
                    if open_break is None:
                        open_break = moment
                elif kind == "lunch_ended" and open_break is not None:
                    breaks.append({"start_at": open_break.isoformat(), "end_at": moment.isoformat(), "duration_seconds": _seconds(open_break, moment), "ongoing": False})
                    open_break = None
                elif kind == "workday_ended" and open_break is not None:
                    breaks.append({"start_at": open_break.isoformat(), "end_at": moment.isoformat(), "duration_seconds": _seconds(open_break, moment), "ongoing": False})
                    open_break = None
            if open_break is not None:
                interval_end = current if day == current.date() else datetime.combine(day + timedelta(days=1), datetime.min.time(), timezone)
                breaks.append({"start_at": open_break.isoformat(), "end_at": None, "duration_seconds": _seconds(open_break, interval_end), "ongoing": day == current.date()})
            break_seconds = sum(item["duration_seconds"] for item in breaks)
            effective_end = shift_end or (current if shift_start and day == current.date() else None)
            work_seconds = (
                max(0, _seconds(shift_start, effective_end) - break_seconds)
                if shift_start and effective_end else None
            )
            items.append({
                "manager_id": manager.get("id"),
                "manager_name": str(manager.get("full_name") or "Менеджер"),
                "date": day.isoformat(),
                "start_at": shift_start.isoformat() if shift_start else None,
                "end_at": shift_end.isoformat() if shift_end else None,
                "breaks": breaks,
                "break_seconds": break_seconds,
                "work_seconds": work_seconds,
                "status": str(manager.get("presence_status") or "unknown") if day == current.date() else ("offline" if shift_end else "unknown"),
            })
    items.sort(key=lambda item: (item["date"], item["manager_name"]), reverse=True)
    return {
        "start": start.isoformat(),
        "end": end.isoformat(),
        "items": items,
        "manager_count": len(raw.get("managers") or []),
        "checked_at": current.isoformat(),
    }


async def collect_worktime(start: date, end: date) -> dict:
    settings = get_settings()
    base_url = settings.leadflow_presence_url.rsplit("/", 1)[0]
    async with httpx.AsyncClient(timeout=30.0) as client:
        response = await client.get(
            f"{base_url}/internal-presence-events",
            params={"start": start.isoformat(), "end": end.isoformat()},
            headers={"X-Internal-Token": settings.leadflow_presence_token},
        )
        response.raise_for_status()
    return build_worktime_payload(
        response.json(), start, end, timezone_name=settings.daily_stats_timezone,
    )


def _time_text(value: str | None) -> str:
    parsed = _parse_datetime(value)
    return parsed.strftime("%H:%M") if parsed else ""


def _duration_text(seconds: int | None) -> str:
    if seconds is None:
        return ""
    hours, remainder = divmod(seconds, 3600)
    minutes = remainder // 60
    return f"{hours:02d}:{minutes:02d}"


def build_worktime_workbook(data: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Табель"
    headers = ["Дата", "Менеджер", "Начало дня", "Конец дня", "Перерывы", "Всего перерывов", "Рабочее время"]
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for item in data["items"]:
        breaks = "; ".join(
            f"{_time_text(entry['start_at'])}–{_time_text(entry['end_at']) or 'сейчас'}"
            for entry in item["breaks"]
        )
        sheet.append([
            datetime.fromisoformat(item["date"]).date(),
            item["manager_name"],
            _time_text(item.get("start_at")),
            _time_text(item.get("end_at")),
            breaks,
            _duration_text(item.get("break_seconds")),
            _duration_text(item.get("work_seconds")),
        ])
        sheet.cell(sheet.max_row, 1).number_format = "DD.MM.YYYY"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate((14, 28, 14, 14, 34, 18, 18), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
