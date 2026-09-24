"""Live Bitrix period audit and source-by-stage Excel reports."""
from __future__ import annotations

import asyncio
import logging
from collections import Counter
from datetime import date, datetime, time, timedelta
from io import BytesIO
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import outcomes, stats
from app.bitrix_client import BitrixClient
from app.config import get_settings
from app.telegram_client import send_telegram_document, send_telegram_message

logger = logging.getLogger(__name__)
_running: dict[int, asyncio.Task] = {}


def _iso_bounds(start: date, end: date, timezone_name: str) -> tuple[str, str]:
    timezone = ZoneInfo(timezone_name)
    start_at = datetime.combine(start, time.min, timezone)
    end_at = datetime.combine(end + timedelta(days=1), time.min, timezone)
    return start_at.isoformat(), end_at.isoformat()


def _stage_map(items: list[dict]) -> dict[str, str]:
    return {
        str(item.get("STATUS_ID") or ""): str(item.get("NAME") or item.get("STATUS_ID") or "Без стадии")
        for item in items
    }


def _after_or_same(item: dict, lead: dict) -> bool:
    item_date = stats._as_datetime(item.get("DATE_CREATE"))
    lead_date = stats._as_datetime(lead.get("DATE_CREATE"))
    return not item_date or not lead_date or item_date >= lead_date


async def collect_live_period(
    start: date,
    end: date,
    *,
    client: BitrixClient | None = None,
) -> dict:
    """Build a fresh lead cohort from Bitrix, independent of local tracking history."""
    settings = get_settings()
    client = client or BitrixClient()
    start_iso, end_iso = _iso_bounds(start, end, settings.daily_stats_timezone)
    contract_field = settings.contract_source_url_field

    leads, sources, lead_statuses, sale_stages, contract_stages = await asyncio.gather(
        client.get_leads_created_between_full(start_iso, end_iso),
        client.get_sources(),
        client.get_lead_statuses(),
        client.get_deal_stages(settings.track_deal_category_id),
        client.get_deal_stages(settings.accompaniment_deal_category_id),
    )
    lead_ids = [str(lead["ID"]) for lead in leads]
    contact_ids = [str(lead.get("CONTACT_ID") or "") for lead in leads]
    explicit_deals, contact_deals, contract_candidates = await asyncio.gather(
        client.get_deals_by_lead_ids(lead_ids, settings.track_deal_category_id),
        client.get_deals_by_contact_ids(contact_ids, settings.track_deal_category_id),
        client.get_deals_created_since(
            settings.accompaniment_deal_category_id,
            start_iso,
            extra_fields=[contract_field],
        ),
    )

    source_names = _stage_map(sources)
    lead_names = _stage_map(lead_statuses)
    sale_names = _stage_map(sale_stages)
    contract_names = _stage_map(contract_stages)
    sales_by_id = {str(deal["ID"]): deal for deal in [*explicit_deals, *contact_deals]}
    lead_id_set = set(lead_ids)
    related_by_lead: dict[str, list[dict]] = {}
    for deal in sales_by_id.values():
        lead_id = str(deal.get("LEAD_ID") or "")
        if lead_id in lead_id_set:
            related_by_lead.setdefault(lead_id, []).append(deal)

    leads_by_contact: dict[str, list[dict]] = {}
    deals_by_contact: dict[str, list[dict]] = {}
    for lead in leads:
        contact_id = str(lead.get("CONTACT_ID") or "")
        if contact_id and contact_id != "0":
            leads_by_contact.setdefault(contact_id, []).append(lead)
    for deal in sales_by_id.values():
        contact_id = str(deal.get("CONTACT_ID") or "")
        if contact_id and contact_id != "0":
            deals_by_contact.setdefault(contact_id, []).append(deal)
    for contact_id, contact_leads in leads_by_contact.items():
        if len(contact_leads) != 1:
            continue
        lead = contact_leads[0]
        lead_id = str(lead["ID"])
        if not related_by_lead.get(lead_id):
            related_by_lead[lead_id] = [
                deal for deal in deals_by_contact.get(contact_id, []) if _after_or_same(deal, lead)
            ]

    contracts_by_source = outcomes._match_contracts(
        contract_candidates, list(sales_by_id.values()), contract_field,
    )
    stage_order = [
        *(f"Лид · {lead_names[str(item.get('STATUS_ID') or '')]}" for item in lead_statuses),
        *(f"Сделка · {sale_names[str(item.get('STATUS_ID') or '')]}" for item in sale_stages),
        *(f"Договор · {contract_names[str(item.get('STATUS_ID') or '')]}" for item in contract_stages),
    ]
    details: list[dict] = []
    for lead in leads:
        lead_id = str(lead["ID"])
        deal = outcomes._pick_related_deal(
            related_by_lead.get(lead_id, []), None, contracts_by_source,
        )
        contract = contracts_by_source.get(str((deal or {}).get("ID") or ""))
        lead_stage_id = str(lead.get("STATUS_ID") or "")
        if contract:
            current_type = "Договор"
            current_stage = contract_names.get(str(contract.get("STAGE_ID") or ""), str(contract.get("STAGE_ID") or "Без стадии"))
        elif deal:
            current_type = "Сделка"
            current_stage = sale_names.get(str(deal.get("STAGE_ID") or ""), str(deal.get("STAGE_ID") or "Без стадии"))
        else:
            current_type = "Лид"
            current_stage = lead_names.get(lead_stage_id, lead_stage_id or "Без стадии")
        source_id = str(lead.get("SOURCE_ID") or "")
        details.append({
            "lead_id": lead_id,
            "created_date": str(lead.get("DATE_CREATE") or "")[:10],
            "source_id": source_id,
            "source_name": source_names.get(source_id, source_id or "Без источника"),
            "lead_stage": lead_names.get(lead_stage_id, lead_stage_id or "Без стадии"),
            "current_type": current_type,
            "current_stage": current_stage,
            "current_label": f"{current_type} · {current_stage}",
            "deal_id": str((deal or {}).get("ID") or ""),
            "deal_stage": sale_names.get(str((deal or {}).get("STAGE_ID") or ""), str((deal or {}).get("STAGE_ID") or "")),
            "contract_id": str((contract or {}).get("ID") or ""),
            "contract_stage": contract_names.get(
                str((contract or {}).get("STAGE_ID") or ""), str((contract or {}).get("STAGE_ID") or ""),
            ),
            "converted": bool(deal),
            "contract": bool(contract),
        })

    observed = {row["current_label"] for row in details}
    ordered_observed = list(dict.fromkeys(label for label in stage_order if label in observed))
    ordered_observed.extend(sorted(observed - set(ordered_observed)))
    return {
        "start": start,
        "end": end,
        "details": details,
        "stage_columns": ordered_observed,
        "checked_at": datetime.now(ZoneInfo(settings.daily_stats_timezone)),
        "portal_domain": urlparse(settings.bitrix_webhook_url).netloc,
    }


def _excel_text(value: object) -> str:
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _style_header(row) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in row:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_workbook(data: dict) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    stages = data["stage_columns"]
    headers = [
        "Источник", *stages, "Перешли в сделку", "Договоры", "Всего лидов",
        "Конверсия в сделку", "Конверсия в договор",
    ]
    summary.append(headers)
    _style_header(summary[1])

    grouped: dict[str, list[dict]] = {}
    for item in data["details"]:
        grouped.setdefault(item["source_name"], []).append(item)
    for source, items in sorted(grouped.items(), key=lambda pair: (-len(pair[1]), pair[0])):
        counts = Counter(item["current_label"] for item in items)
        converted = sum(item["converted"] for item in items)
        contracts = sum(item["contract"] for item in items)
        total = len(items)
        summary.append([
            _excel_text(source), *(counts.get(stage, 0) for stage in stages),
            converted, contracts, total, converted / total if total else 0, contracts / total if total else 0,
        ])
    totals = data["details"]
    total_counts = Counter(item["current_label"] for item in totals)
    converted_total = sum(item["converted"] for item in totals)
    contract_total = sum(item["contract"] for item in totals)
    grand_total = len(totals)
    summary.append([
        "ИТОГО", *(total_counts.get(stage, 0) for stage in stages), converted_total,
        contract_total, grand_total, converted_total / grand_total if grand_total else 0,
        contract_total / grand_total if grand_total else 0,
    ])
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    summary.freeze_panes = "B2"
    summary.auto_filter.ref = summary.dimensions
    for column in range(len(headers) - 1, len(headers) + 1):
        for row in range(2, summary.max_row + 1):
            summary.cell(row, column).number_format = "0.0%"
    for index, header in enumerate(headers, start=1):
        summary.column_dimensions[get_column_letter(index)].width = min(42, max(12, len(header) + 2))

    detail = workbook.create_sheet("Детализация")
    detail_headers = [
        "Дата создания", "ID лида", "Источник", "Текущий объект", "Текущая стадия",
        "ID сделки", "Стадия сделки", "ID договора", "Стадия сопровождения",
    ]
    detail.append(detail_headers)
    _style_header(detail[1])
    portal = data["portal_domain"]
    for item in data["details"]:
        detail.append([
            item["created_date"], item["lead_id"], _excel_text(item["source_name"]),
            item["current_type"], _excel_text(item["current_stage"]), item["deal_id"],
            _excel_text(item["deal_stage"]), item["contract_id"], _excel_text(item["contract_stage"]),
        ])
        row = detail.max_row
        detail.cell(row, 2).hyperlink = f"https://{portal}/crm/lead/details/{item['lead_id']}/"
        detail.cell(row, 2).style = "Hyperlink"
        if item["deal_id"]:
            detail.cell(row, 6).hyperlink = f"https://{portal}/crm/deal/details/{item['deal_id']}/"
            detail.cell(row, 6).style = "Hyperlink"
        if item["contract_id"]:
            detail.cell(row, 8).hyperlink = f"https://{portal}/crm/deal/details/{item['contract_id']}/"
            detail.cell(row, 8).style = "Hyperlink"
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    widths = [15, 12, 28, 18, 32, 12, 30, 14, 32]
    for index, width in enumerate(widths, start=1):
        detail.column_dimensions[get_column_letter(index)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_caption(data: dict) -> str:
    details = data["details"]
    converted = sum(item["converted"] for item in details)
    contracts = sum(item["contract"] for item in details)
    total = len(details)
    converted_rate = converted / total * 100 if total else 0
    contract_rate = contracts / total * 100 if total else 0
    return (
        f"📊 <b>{data['start']:%d.%m.%Y}–{data['end']:%d.%m.%Y}</b>\n"
        f"Лидов: <b>{total}</b>\n"
        f"Перешли в сделку: <b>{converted}</b> ({converted_rate:.1f}%)\n"
        f"Договоры: <b>{contracts}</b> ({contract_rate:.1f}%)\n"
        f"Источников: <b>{len({item['source_name'] for item in details})}</b>\n"
        f"<i>Сверено с Bitrix: {data['checked_at']:%d.%m.%Y %H:%M}</i>"
    )


async def deliver_period_report(start: date, end: date, chat_id: int) -> None:
    try:
        data = await collect_live_period(start, end)
        content = build_workbook(data)
        filename = f"leads_{start:%d.%m.%Y}-{end:%d.%m.%Y}.xlsx"
        await send_telegram_document(
            content, filename, chat_id=chat_id, caption=build_caption(data),
        )
    except Exception:
        logger.exception("Live period report failed chat_id=%s start=%s end=%s", chat_id, start, end)
        try:
            await send_telegram_message(
                "Не удалось сверить отчёт с Bitrix. Ошибка записана в журнал, попробуйте позже.",
                chat_id=chat_id,
            )
        except Exception:
            logger.exception("Failed to notify chat_id=%s about period report error", chat_id)


def enqueue_period_report(start: date, end: date, chat_id: int) -> bool:
    existing = _running.get(chat_id)
    if existing and not existing.done():
        return False
    task = asyncio.create_task(deliver_period_report(start, end, chat_id))
    _running[chat_id] = task
    task.add_done_callback(lambda finished, key=chat_id: _running.pop(key, None))
    return True
