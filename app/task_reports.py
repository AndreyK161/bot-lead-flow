"""Standardized sales task report with manual/automation separation."""
from __future__ import annotations

import asyncio
from datetime import date, datetime, time, timedelta
from io import BytesIO
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import store
from app.bitrix_client import BitrixClient
from app.config import get_settings
from app.sales_team import assignable_sales_users

TASK_STATUS_NAMES = {
    "1": "Новая", "2": "Ждёт выполнения", "3": "Выполняется",
    "4": "Ждёт контроля", "5": "Завершена", "6": "Отложена", "7": "Отклонена",
}
CALL_WORDS = ("позвон", "прозвон", "связаться", "звонок")


def _value(item: dict, *keys: str):
    for key in keys:
        if item.get(key) is not None:
            return item[key]
    return None


def _parse_datetime(value: object) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None


def _phone(value: object) -> str:
    digits = "".join(char for char in str(value or "") if char.isdigit())
    return digits[-10:] if len(digits) >= 10 else digits


def _first_phone(entity: dict) -> str:
    for item in entity.get("PHONE") or []:
        value = str((item or {}).get("VALUE") or "").strip()
        if value:
            return value
    return ""


def _binding(task: dict) -> tuple[str, str]:
    bindings = _value(task, "ufCrmTask", "UF_CRM_TASK") or []
    if isinstance(bindings, str):
        bindings = [bindings]
    for raw in bindings:
        value = str(raw)
        if value.startswith("D_"):
            return "deal", value[2:]
        if value.startswith("CRM_DEAL_"):
            return "deal", value[9:]
    for raw in bindings:
        value = str(raw)
        if value.startswith("L_"):
            return "lead", value[2:]
    return "", ""


def _callback_status(phone: str, responsible_id: str, deadline: datetime | None, settings) -> str:
    if not phone or not deadline:
        return "Нет данных"
    extensions = [
        extension for extension, user_id in getattr(settings, "megafon_bitrix_users", {}).items()
        if str(user_id) == responsible_id
    ]
    if not extensions:
        return "Нет привязки телефонии"
    start = deadline - timedelta(minutes=15)
    end = deadline + timedelta(minutes=15)
    placeholders = ",".join("?" for _ in extensions)
    rows = store.rows(
        f"""SELECT 1 FROM megafon_calls
            WHERE client_phone=? AND manager_extension IN ({placeholders})
              AND direction='out' AND started_at>=? AND started_at<=? LIMIT 1""",
        (_phone(phone), *extensions, start.isoformat(), end.isoformat()),
    )
    if rows:
        return "Вовремя"
    now = datetime.now(deadline.tzinfo or ZoneInfo(settings.daily_stats_timezone))
    return "Ожидается" if now <= end else "Не зафиксирован вовремя"


async def collect_task_report(
    start: date, end: date, *, client: BitrixClient | None = None,
) -> dict:
    settings = get_settings()
    client = client or BitrixClient()
    timezone = ZoneInfo(settings.daily_stats_timezone)
    start_at = datetime.combine(start, time.min, timezone)
    end_at = datetime.combine(end + timedelta(days=1), time.min, timezone)
    users = assignable_sales_users(
        await client.get_department_users(settings.sales_department_id), settings,
    )
    manager_names = {
        str(user["ID"]): " ".join(
            str(user.get(key) or "").strip() for key in ("NAME", "LAST_NAME")
        ).strip() or f"Сотрудник #{user['ID']}"
        for user in users
    }
    tasks = await client.get_tasks_by_deadline(
        start_at.isoformat(), end_at.isoformat(), list(manager_names),
    )
    task_bindings = [_binding(task) for task in tasks]
    lead_ids = [entity_id for entity_type, entity_id in task_bindings if entity_type == "lead"]
    deal_ids = [entity_id for entity_type, entity_id in task_bindings if entity_type == "deal"]
    leads, deals = await asyncio.gather(
        client.get_leads_by_ids(lead_ids), client.get_deals_by_ids(deal_ids),
    )
    contact_ids = [str(deal.get("CONTACT_ID") or "") for deal in deals]
    contacts = await client.get_contacts_full_by_ids(contact_ids)
    leads_by_id = {str(item["ID"]): item for item in leads}
    deals_by_id = {str(item["ID"]): item for item in deals}
    contacts_by_id = {str(item["ID"]): item for item in contacts}
    automation_ids = getattr(settings, "bitrix_automation_user_id_set", {"24"})
    portal = urlparse(settings.bitrix_webhook_url).netloc
    items = []
    for task, (entity_type, entity_id) in zip(tasks, task_bindings):
        responsible_id = str(_value(task, "responsibleId", "RESPONSIBLE_ID") or "")
        creator_id = str(_value(task, "createdBy", "CREATED_BY") or "")
        creator = task.get("creator") or {}
        creator_name = str(creator.get("name") or creator.get("NAME") or f"Сотрудник #{creator_id}")
        title = str(_value(task, "title", "TITLE") or "Без названия")
        task_type = "Звонок" if any(word in title.casefold() for word in CALL_WORDS) else "Задача"
        if entity_type == "lead":
            phone = _first_phone(leads_by_id.get(entity_id, {}))
        else:
            deal = deals_by_id.get(entity_id, {})
            phone = _first_phone(contacts_by_id.get(str(deal.get("CONTACT_ID") or ""), {}))
        deadline = _parse_datetime(_value(task, "deadline", "DEADLINE"))
        origin = "automatic" if creator_id in automation_ids else "manual"
        callback_status = (
            _callback_status(phone, responsible_id, deadline, settings)
            if task_type == "Звонок" and origin == "manual" else "Не применяется"
        )
        task_id = str(_value(task, "id", "ID") or "")
        items.append({
            "task_id": task_id,
            "task_url": f"https://{portal}/company/personal/user/{responsible_id}/tasks/task/view/{task_id}/",
            "title": title,
            "type": task_type,
            "origin": origin,
            "responsible_id": responsible_id,
            "responsible_name": manager_names.get(responsible_id, f"Сотрудник #{responsible_id}"),
            "creator_id": creator_id,
            "creator_name": creator_name,
            "created_at": str(_value(task, "createdDate", "CREATED_DATE") or ""),
            "deadline": str(_value(task, "deadline", "DEADLINE") or ""),
            "closed_at": str(_value(task, "closedDate", "CLOSED_DATE") or ""),
            "status": TASK_STATUS_NAMES.get(str(_value(task, "status", "STATUS") or ""), str(_value(task, "status", "STATUS") or "")),
            "entity_type": entity_type,
            "entity_id": entity_id,
            "deal_id": entity_id if entity_type == "deal" else "",
            "phone": phone,
            "callback_status": callback_status,
        })
    return {
        "start": start.isoformat(), "end": end.isoformat(), "items": items,
        "totals": {
            "all": len(items),
            "manual": sum(item["origin"] == "manual" for item in items),
            "automatic": sum(item["origin"] == "automatic" for item in items),
            "calls": sum(item["type"] == "Звонок" for item in items),
            "on_time": sum(item["callback_status"] == "Вовремя" for item in items),
        },
    }


def _excel_text(value: object) -> str:
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def build_task_workbook(data: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Дела и задачи"
    headers = [
        "Тип", "Название", "Ответственный", "Постановщик", "Источник постановки",
        "Крайний срок", "Дата выполнения", "Статус", "ID сделки", "Телефон клиента",
        "Перезвон ±15 минут", "ID задачи",
    ]
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill; cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for item in data["items"]:
        sheet.append([
            item["type"], _excel_text(item["title"]), item["responsible_name"],
            item["creator_name"], "Робот" if item["origin"] == "automatic" else "Человек",
            item["deadline"], item["closed_at"], item["status"], item["deal_id"],
            _excel_text(item["phone"]), item["callback_status"], item["task_id"],
        ])
        row = sheet.max_row
        sheet.cell(row, 10).number_format = "@"
        sheet.cell(row, 12).hyperlink = item["task_url"]
        sheet.cell(row, 12).style = "Hyperlink"
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate((12, 34, 25, 25, 20, 22, 22, 20, 14, 20, 22, 14), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO(); workbook.save(buffer); return buffer.getvalue()
