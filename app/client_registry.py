"""Current category-2 client registry and its Excel export."""
from __future__ import annotations

import asyncio
import re
from datetime import datetime
from io import BytesIO
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.bitrix_client import BitrixClient
from app.config import get_settings

SOURCE_DEAL_RE = re.compile(r"/dogovor/(\d+)(?:/|\?|$)", re.IGNORECASE)


def _contact_name(contact: dict, deal: dict) -> str:
    parts = [
        str(contact.get(key) or "").strip()
        for key in ("LAST_NAME", "NAME", "SECOND_NAME")
    ]
    full_name = " ".join(part for part in parts if part)
    return full_name or str(deal.get("TITLE") or f"Сделка №{deal.get('ID')}")


def _contact_phone(contact: dict) -> str:
    phones = contact.get("PHONE") or []
    if isinstance(phones, list):
        for phone in phones:
            value = str((phone or {}).get("VALUE") or "").strip()
            if value:
                return value
    return ""


def _as_datetime(value: object) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None


async def collect_clients(*, client: BitrixClient | None = None) -> dict:
    """Load all active/successful accompaniment deals from Bitrix."""
    settings = get_settings()
    client = client or BitrixClient()
    category_id = settings.accompaniment_deal_category_id
    contract_field = settings.deal_contract_date_field
    source_field = getattr(settings, "contract_source_url_field", "UF_CRM_1775217002")
    deals, stages = await asyncio.gather(
        client.get_category_deals(category_id, extra_fields=[contract_field, source_field]),
        client.get_deal_stages(category_id),
    )
    stage_by_id = {str(item.get("STATUS_ID") or ""): item for item in stages}
    included = [
        deal for deal in deals
        if str(stage_by_id.get(str(deal.get("STAGE_ID") or ""), {}).get("SEMANTICS") or "P") != "F"
    ]
    contact_ids = [str(deal.get("CONTACT_ID") or "") for deal in included]
    contacts = await client.get_contacts_full_by_ids(contact_ids)
    contact_by_id = {str(contact.get("ID")): contact for contact in contacts}
    source_ids = []
    source_id_by_contract: dict[str, str] = {}
    for deal in included:
        match = SOURCE_DEAL_RE.search(str(deal.get(source_field) or ""))
        if match:
            source_id_by_contract[str(deal.get("ID") or "")] = match.group(1)
            source_ids.append(match.group(1))
    source_deals = await client.get_deals_by_ids(source_ids) if source_ids else []
    source_by_id = {str(deal.get("ID")): deal for deal in source_deals}
    source_contact_ids = [str(deal.get("CONTACT_ID") or "") for deal in source_deals]
    source_contacts = await client.get_contacts_full_by_ids(source_contact_ids) if source_contact_ids else []
    source_contact_by_id = {str(contact.get("ID")): contact for contact in source_contacts}
    source_lead_ids = [str(deal.get("LEAD_ID") or "") for deal in source_deals if deal.get("LEAD_ID")]
    source_leads = await client.get_leads_by_ids(source_lead_ids) if source_lead_ids else []
    source_lead_by_id = {str(lead.get("ID")): lead for lead in source_leads}
    portal_domain = urlparse(settings.bitrix_webhook_url).netloc

    items = []
    for deal in included:
        deal_id = str(deal.get("ID") or "")
        contact = contact_by_id.get(str(deal.get("CONTACT_ID") or ""), {})
        source_deal = source_by_id.get(source_id_by_contract.get(deal_id, ""), {})
        source_contact = source_contact_by_id.get(str(source_deal.get("CONTACT_ID") or ""), {})
        source_lead = source_lead_by_id.get(str(source_deal.get("LEAD_ID") or ""), {})
        stage = stage_by_id.get(str(deal.get("STAGE_ID") or ""), {})
        contract_date = deal.get(contract_field) or deal.get("DATE_CREATE") or ""
        semantics = str(stage.get("SEMANTICS") or "P")
        items.append({
            "deal_id": deal_id,
            "deal_url": f"https://{portal_domain}/crm/deal/details/{deal_id}/",
            "name": _contact_name(contact or source_contact, deal),
            "phone": _contact_phone(contact) or _contact_phone(source_contact) or _contact_phone(source_lead),
            "stage": str(stage.get("NAME") or deal.get("STAGE_ID") or "Без стадии"),
            "contract_date": str(contract_date),
            "status": "completed" if semantics == "S" else "active",
        })
    items.sort(
        key=lambda item: (str(item["contract_date"]), int(item["deal_id"] or 0)),
        reverse=True,
    )
    return {
        "items": items,
        "totals": {
            "all": len(items),
            "active": sum(item["status"] == "active" for item in items),
            "completed": sum(item["status"] == "completed" for item in items),
            "without_phone": sum(not item["phone"] for item in items),
        },
        "checked_at": datetime.now(ZoneInfo(settings.daily_stats_timezone)).isoformat(),
    }


def _excel_text(value: object) -> str:
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def build_clients_workbook(data: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Клиенты"
    headers = ["ФИО", "Телефон", "Стадия сопровождения", "Дата договора", "Статус", "Сделка"]
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for item in data["items"]:
        parsed_date = _as_datetime(item.get("contract_date"))
        if parsed_date and parsed_date.tzinfo:
            parsed_date = parsed_date.replace(tzinfo=None)
        sheet.append([
            _excel_text(item.get("name")),
            _excel_text(item.get("phone")),
            _excel_text(item.get("stage")),
            parsed_date or _excel_text(item.get("contract_date")),
            "Завершён" if item.get("status") == "completed" else "Активен",
            item.get("deal_id"),
        ])
        row = sheet.max_row
        sheet.cell(row, 2).number_format = "@"
        if parsed_date:
            sheet.cell(row, 4).number_format = "DD.MM.YYYY HH:MM"
        sheet.cell(row, 6).hyperlink = item.get("deal_url")
        sheet.cell(row, 6).style = "Hyperlink"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate((32, 20, 32, 20, 14, 14), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
