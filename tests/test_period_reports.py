import tempfile
import unittest
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from app import period_reports


TEST_DIRECTOR_ID = 1297686797


class LivePeriodReportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.settings = SimpleNamespace(
            daily_stats_timezone="Europe/Moscow",
            contract_source_url_field="UF_CONTRACT_URL",
            track_deal_category_id="0",
            accompaniment_deal_category_id="2",
            bitrix_webhook_url="https://example.bitrix24.ru/rest/1/token/",
            director_user_id_set={TEST_DIRECTOR_ID},
        )
        self.settings_patch = patch("app.period_reports.get_settings", return_value=self.settings)
        self.settings_patch.start()

    async def asyncTearDown(self):
        self.settings_patch.stop()

    async def test_live_collection_finds_contract_and_keeps_original_source(self):
        client = AsyncMock()
        client.get_leads_created_between_full.return_value = [
            {
                "ID": "10", "SOURCE_ID": "WEB", "STATUS_ID": "CONVERTED", "CONTACT_ID": "500",
                "DATE_CREATE": "2026-09-02T10:00:00+03:00",
            },
            {
                "ID": "11", "SOURCE_ID": "TG", "STATUS_ID": "IN_PROCESS", "CONTACT_ID": "501",
                "DATE_CREATE": "2026-09-03T10:00:00+03:00",
            },
        ]
        client.get_sources.return_value = [
            {"STATUS_ID": "WEB", "NAME": "Сайт"},
            {"STATUS_ID": "TG", "NAME": "Телеграм"},
            {"STATUS_ID": "VK", "NAME": "ВКонтакте"},
        ]
        client.get_lead_statuses.return_value = [
            {"STATUS_ID": "IN_PROCESS", "NAME": "Недозвон"},
            {"STATUS_ID": "CONVERTED", "NAME": "Качественный лид"},
        ]
        client.get_deal_stages.side_effect = [
            [
                {"STATUS_ID": "NEW", "NAME": "Не обработан"},
                {"STATUS_ID": "WON", "NAME": "Сделка успешна"},
            ],
            [{"STATUS_ID": "C2:NEW", "NAME": "Новый договор"}],
        ]
        sale = {
            "ID": "20", "LEAD_ID": "10", "CONTACT_ID": "500", "CATEGORY_ID": "0",
            "STAGE_ID": "WON", "DATE_CREATE": "2026-09-04T10:00:00+03:00",
        }
        client.get_deals_by_lead_ids.return_value = [sale]
        client.get_deals_by_contact_ids.return_value = [sale]
        client.get_deals_created_between_full.return_value = [sale, {
            "ID": "21", "LEAD_ID": None, "CONTACT_ID": "900", "CATEGORY_ID": "0",
            "SOURCE_ID": "VK", "STAGE_ID": "NEW", "DATE_CREATE": "2026-09-06T10:00:00+03:00",
        }]
        client.get_deals_created_since.return_value = [{
            "ID": "30", "CONTACT_ID": "500", "CATEGORY_ID": "2", "STAGE_ID": "C2:NEW",
            "DATE_CREATE": "2026-09-05T10:00:00+03:00",
            "UF_CONTRACT_URL": "https://example.test/dogovor/20/?token=hidden",
        }]

        data = await period_reports.collect_live_period(
            date(2026, 9, 1), date(2026, 9, 23), client=client,
        )

        self.assertEqual(len(data["leads"]), 2)
        self.assertEqual(len(data["deals"]), 2)
        self.assertEqual(data["unique_total"], 3)
        self.assertEqual(data["direct_deal_count"], 1)
        contract = next(item for item in data["leads"] if item["lead_id"] == "10")
        self.assertEqual(contract["source_name"], "Сайт")
        self.assertEqual(contract["deal_id"], "20")
        self.assertEqual(contract["contract_id"], "30")
        self.assertEqual(contract["result_stage"], "Сконвертирован")
        self.assertIn("Недозвон", data["lead_stage_columns"])

        content = period_reports.build_workbook(data)
        workbook = load_workbook(BytesIO(content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Итоги", "Лиды", "Сделки", "Детализация"])
        lead_sheet = workbook["Лиды"]
        lead_headers = [cell.value for cell in lead_sheet[1]]
        self.assertEqual(lead_headers[-1], "Всего лидов")
        site_row = next(row for row in lead_sheet.iter_rows(values_only=True) if row[0] == "Сайт")
        self.assertEqual(site_row[lead_headers.index("Сконвертирован")], 1)
        deal_sheet = workbook["Сделки"]
        deal_headers = [cell.value for cell in deal_sheet[1]]
        vk_row = next(row for row in deal_sheet.iter_rows(values_only=True) if row[0] == "ВКонтакте")
        self.assertEqual(vk_row[deal_headers.index("Не обработан")], 1)
        self.assertEqual(deal_headers[-2:], ["Договоры", "Всего сделок"])
        detail = workbook["Детализация"]
        self.assertEqual(detail.max_row, 5)
        self.assertIsNotNone(detail["C2"].hyperlink)

    def test_excel_escapes_formula_like_source_names(self):
        data = {
            "start": date(2026, 9, 1),
            "end": date(2026, 9, 1),
            "checked_at": None,
            "portal_domain": "example.bitrix24.ru",
            "lead_stage_columns": ["Не обработан"],
            "deal_stage_columns": [],
            "direct_deal_count": 0,
            "unique_total": 1,
            "deals": [],
            "leads": [{
                "lead_id": "1", "created_date": "2026-09-01", "source_name": "=FORMULA",
                "result_stage": "Не обработан", "deal_id": "", "deal_stage": "",
                "contract_id": "", "contract_stage": "", "converted": False, "contract": False,
            }],
        }
        workbook = load_workbook(BytesIO(period_reports.build_workbook(data)), data_only=False)
        self.assertEqual(workbook["Лиды"]["A2"].value, "'=FORMULA")


if __name__ == "__main__":
    unittest.main()
