import tempfile
import unittest
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from app import store, telephony_exports


class FakeClient:
    async def get_user_name(self, user_id):
        return "Яна Азизова"


class TelephonyExportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.settings = SimpleNamespace(
            database_path=self.tmp.name + "/db.sqlite", database_url="",
            daily_stats_timezone="Europe/Moscow", megafon_bitrix_users={"706": "110"},
        )
        self.settings_patch = patch("app.telephony_exports.get_settings", return_value=self.settings)
        self.store_patch = patch("app.store.get_settings", return_value=self.settings)
        self.settings_patch.start(); self.store_patch.start()
        store.execute(
            """INSERT INTO megafon_calls
               (call_id,manager_extension,direction,client_phone,started_at,duration_seconds,
                wait_seconds,call_status,processing_status)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            ("call-1", "706", "out", "9990001122", "2026-10-19T09:30:00+00:00", 125, 9, "Success", "not_applicable"),
        )

    async def asyncTearDown(self):
        self.settings_patch.stop(); self.store_patch.stop(); self.tmp.cleanup()

    async def test_fixed_telephony_sheets(self):
        data = await telephony_exports.collect_telephony_export(
            date(2026, 10, 19), date(2026, 10, 19), client=FakeClient(),
        )
        self.assertEqual(data["items"][0]["employee"], "Яна Азизова")
        self.assertEqual(data["items"][0]["wait_seconds"], 9)
        self.assertTrue(data["items"][0]["answered"])
        workbook = load_workbook(BytesIO(telephony_exports.build_telephony_workbook(data)))
        self.assertEqual(workbook.sheetnames, ["Телефония АТС", "Мои звонки"])
        self.assertEqual([cell.value for cell in workbook["Телефония АТС"][1]], [
            "Тип звонка", "Клиент", "Сотрудник", "Дата", "Время", "Ожидание", "Длительность",
        ])
        self.assertEqual([cell.value for cell in workbook["Мои звонки"][1]], [
            "Дата", "Сотрудник", "Номер", "Направление", "Отвечен", "Длительность", "Дозвон",
        ])
        self.assertEqual(workbook["Телефония АТС"]["F2"].value, "00:00:09")


if __name__ == "__main__":
    unittest.main()
